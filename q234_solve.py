#!/usr/bin/env python3
"""2026 高教社杯 C 题：问题 2、3、4 的可复现求解程序。

直接运行：
    python q234_solve.py

快速检查（只计算少量日期，不写正式模板）：
    python q234_solve.py --smoke

依赖安装：
    python -m pip install numpy pandas scipy matplotlib openpyxl pillow

建模口径：
1. 功率乘 1/6 h 转为每 10 分钟电量，所有平衡式统一使用 kWh。
2. 问题 2 使用严格因果的滚动净负荷分位数预测；0:00 后不修改计划购电量。
3. 问题 3 使用附件 3 的 0/6/12/18 点光伏预报，历史时段冻结，只重算未来时段。
4. 下调时取消原购电但仍支付 50% 违约费，因此相对原计划成本的变化为 -0.5p×下调量；
   上调超出原计划部分按 1.5p 计费。该解释与题面“违约电价 50%”一致。
5. 问题 4 假定附件 4 的当日电价在 0:00 已知。若需要“只知历史价格”，应另接价格预测器。
6. 问题 2 至 4 统一采用含充放电二元互斥的 MILP，并用两层词典序优化分离题目费用与吞吐量。
7. 储能状态按实际执行结果跨日连续，仅在整个计算区间最后一天回到初始储电量。

程序不会覆盖附件原件，只在 results/、figures/ 中生成结果。
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
import time
from copy import copy
from dataclasses import asdict, dataclass
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "附件"
TEMPLATE_DIR = DATA_DIR / "附件5"
RESULTS_DIR = ROOT / "results"
FIGURES_DIR = ROOT / "figures"


@dataclass(frozen=True)
class Parameters:
    dt_hours: float = 1.0 / 6.0
    capacity_kwh: float = 12000.0
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    soc_initial_kwh: float = 6000.0
    max_power_kw: float = 5000.0
    eta_charge: float = 0.90
    eta_discharge: float = 0.90
    lexicographic_abs_tolerance_yuan: float = 1e-5
    lexicographic_rel_tolerance: float = 1e-9
    emergency_multiplier: float = 5.0
    up_adjust_multiplier: float = 1.5
    down_cancel_refund_multiplier: float = 0.5
    lookback_days: int = 56
    tolerance: float = 1e-6

    @property
    def max_interval_energy_kwh(self) -> float:
        return self.max_power_kw * self.dt_hours


@dataclass(frozen=True)
class DynamicQuantileParameters:
    """正式日前策略使用的因果滚动净负荷分位数选参配置。"""

    training_days: int = 60
    validation_days: int = 14
    similar_days: int = 20
    candidates: tuple[float, ...] = tuple(np.round(np.arange(0.50, 0.951, 0.05), 2))
    cvar_level: float = 0.95
    risk_weight: float = 0.20
    stability_weight: float = 0.02
    default_alpha: float = 0.80


@dataclass
class InputData:
    dates: pd.DatetimeIndex
    time_labels: list[str]
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    fixed_price: np.ndarray
    variable_price: np.ndarray
    official_pv_kw: np.ndarray
    fallback_load_kwh: np.ndarray
    fallback_pv_kwh: np.ndarray


@dataclass
class OptimizationResult:
    grid: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    spill: np.ndarray
    soc_end: np.ndarray
    objective: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="问题2至4：微网因果预测与滚动调度")
    parser.add_argument("--smoke", action="store_true", help="只计算 1 月 1 日至 2 月 3 日并做约束校验")
    parser.add_argument("--eta", type=float, default=0.90, help="单程效率，必须在 (0,1] 范围内")
    return parser.parse_args()


def _numeric_matrix(path: Path, sheet_name: str) -> tuple[pd.DatetimeIndex, list[str], np.ndarray]:
    raw = pd.read_excel(path, sheet_name=sheet_name, header=0)
    if raw.shape != (365, 145):
        raise ValueError(f"{path.name}/{sheet_name} 应为 365×145，实际为 {raw.shape}")
    dates = pd.DatetimeIndex(pd.to_datetime(raw.iloc[:, 0], errors="raise"))
    if dates.duplicated().any() or not dates.is_monotonic_increasing:
        raise ValueError(f"{path.name}/{sheet_name} 的日期必须唯一且递增")
    values = raw.iloc[:, 1:].apply(pd.to_numeric, errors="raise").to_numpy(float)
    if not np.isfinite(values).all() or (values < 0).any():
        raise ValueError(f"{path.name}/{sheet_name} 含缺失、无穷或负数")
    return dates, [str(x) for x in raw.columns[1:]], values


def _time_start_minutes(value) -> int:
    """把附件列名解释为时段起点；末列 0:00+1 为次日 0:00。"""
    text = str(value).strip()
    if "+1" in text:
        return 1440
    match = re.search(r"(\d{1,2}):(\d{2})", text)
    if not match:
        raise ValueError(f"无法识别时间标签：{value}")
    return int(match.group(1)) * 60 + int(match.group(2))


def _issue_index(time_labels: list[str], hour: int) -> int:
    """返回整点开始的首个未执行时段；6:00 对应 6:00-6:10，而不是 6:10-6:20。"""
    if hour == 0:
        return 0
    target = hour * 60
    matches = [i for i, label in enumerate(time_labels) if _time_start_minutes(label) == target]
    if len(matches) != 1:
        raise ValueError(f"时间轴中无法唯一定位 {hour}:00：{matches}")
    return matches[0]


def load_inputs(params: Parameters) -> InputData:
    """数据输入：核对四个附件，统一把 kW 转成 kWh。"""
    dates, labels, load_kw = _numeric_matrix(DATA_DIR / "附件2.xlsx", "小区负载")
    dates_pv, labels_pv, pv_kw = _numeric_matrix(DATA_DIR / "附件2.xlsx", "光伏发电实际功率")
    dates_price, labels_price, variable_price = _numeric_matrix(DATA_DIR / "附件4.xlsx", "Sheet1")
    if not dates.equals(dates_pv) or not dates.equals(dates_price):
        raise ValueError("附件2与附件4日期键不一致")
    if labels != labels_pv or labels != labels_price:
        raise ValueError("附件2与附件4的 144 个时刻键不一致")
    if [_time_start_minutes(x) for x in labels] != list(range(10, 1450, 10)):
        raise ValueError("附件2/4时间必须依次表示 00:10-00:20 至次日 00:00-00:10")

    fixed = pd.read_excel(DATA_DIR / "附件1.xlsx", sheet_name=0)
    expected = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    if list(fixed.columns) != expected or len(fixed) != 144:
        raise ValueError("附件1表头或时段数异常")
    fixed_num = fixed[expected[1:]].apply(pd.to_numeric, errors="raise")

    forecast = pd.read_excel(DATA_DIR / "附件3.xlsx", sheet_name=0)
    if forecast.shape != (1460, 26):
        raise ValueError(f"附件3应有 1460 条预报，实际为 {forecast.shape}")
    forecast["日期"] = pd.to_datetime(forecast["日期"].ffill(), errors="raise")
    issue_map = {"0:00": 0, "6:00": 1, "12:00": 2, "18:00": 3}
    official = np.full((365, 4, 24), np.nan)
    date_to_i = {d.normalize(): i for i, d in enumerate(dates)}
    for _, row in forecast.iterrows():
        day_i = date_to_i[pd.Timestamp(row["日期"]).normalize()]
        issue_text = str(row["预报时刻"]).strip()
        if issue_text not in issue_map:
            raise ValueError(f"附件3出现未知预报时刻：{issue_text}")
        official[day_i, issue_map[issue_text], :] = pd.to_numeric(row.iloc[2:], errors="raise").to_numpy(float)
    if not np.isfinite(official).all() or (official < 0).any():
        raise ValueError("附件3预报存在缺失或负数")

    return InputData(
        dates=dates,
        time_labels=labels,
        load_kwh=load_kw * params.dt_hours,
        pv_kwh=pv_kw * params.dt_hours,
        fixed_price=fixed_num["电价"].to_numpy(float),
        variable_price=variable_price,
        official_pv_kw=official,
        fallback_load_kwh=fixed_num["小区负载"].to_numpy(float) * params.dt_hours,
        fallback_pv_kwh=fixed_num["光伏发电预测功率"].to_numpy(float) * params.dt_hours,
    )


def validate_parameters(p: Parameters) -> None:
    if not (0 < p.eta_charge <= 1 and 0 < p.eta_discharge <= 1):
        raise ValueError("参数 eta_charge、eta_discharge 需在 (0,1] 范围内调试")
    if p.soc_min_kwh >= p.soc_max_kwh or p.max_power_kw <= 0:
        raise ValueError("储能上下界或最大功率参数无效")
    if p.lexicographic_abs_tolerance_yuan < 0 or p.lexicographic_rel_tolerance < 0:
        raise ValueError("词典序优化容差不得为负数")


def causal_candidates(day_i: int, dates: pd.DatetimeIndex, lookback: int) -> np.ndarray:
    """只选择当前日之前的相似工作日，杜绝把未来实际值用于计划。"""
    start = max(0, day_i - lookback)
    all_past = np.arange(start, day_i, dtype=int)
    if len(all_past) == 0:
        return all_past
    same_weekday = all_past[dates[all_past].weekday == dates[day_i].weekday()]
    return same_weekday if len(same_weekday) >= 3 else all_past


def dynamic_similar_days(
    data: InputData,
    day_i: int,
    cfg: DynamicQuantileParameters,
) -> np.ndarray:
    """只在滚动训练窗内选择过去相似日，不使用当前日及未来实际数据。"""
    start = max(0, day_i - cfg.training_days)
    candidates = np.arange(start, day_i, dtype=int)
    if not len(candidates):
        return candidates
    weekday = data.dates[day_i].weekday()
    same_type = candidates[data.dates[candidates].weekday == weekday]
    pool = same_type if len(same_type) >= min(5, cfg.similar_days) else candidates
    recent_start = max(0, day_i - 7)
    reference = np.mean(
        data.load_kwh[recent_start:day_i] - data.pv_kwh[recent_start:day_i], axis=0
    )
    net = data.load_kwh[pool] - data.pv_kwh[pool]
    scale = max(float(np.std(net)), 1.0)
    shape_distance = np.sqrt(np.mean(((net - reference) / scale) ** 2, axis=1))
    recency = (day_i - pool) / max(cfg.training_days, 1)
    score = shape_distance + 0.15 * recency
    return pool[np.argsort(score)[: cfg.similar_days]]


def dynamic_net_forecast(
    data: InputData,
    day_i: int,
    alpha: float,
    cfg: DynamicQuantileParameters,
) -> np.ndarray:
    """按候选分位数直接预测净负荷，避免拆分边际分位数造成口径错配。"""
    indices = dynamic_similar_days(data, day_i, cfg)
    if not len(indices):
        return data.fallback_load_kwh - data.fallback_pv_kwh
    paired_net = data.load_kwh[indices] - data.pv_kwh[indices]
    return np.quantile(paired_net, alpha, axis=0)


def _hourly_sum(values: np.ndarray) -> np.ndarray:
    return np.asarray(values, dtype=float).reshape(24, 6).sum(axis=1)


def dynamic_validation_cost(
    data: InputData,
    day_i: int,
    alpha: float,
    cfg: DynamicQuantileParameters,
    p: Parameters,
    price_mode: str,
) -> float:
    """用已实现历史日的一小时聚合调度成本评价候选分位数。"""
    net_hat = _hourly_sum(dynamic_net_forecast(data, day_i, alpha, cfg))
    price_10min = data.fixed_price if price_mode == "fixed" else data.variable_price[day_i]
    price = np.asarray(price_10min, dtype=float).reshape(24, 6).mean(axis=1)
    hourly_params = Parameters(
        dt_hours=1.0,
        capacity_kwh=p.capacity_kwh,
        soc_min_kwh=p.soc_min_kwh,
        soc_max_kwh=p.soc_max_kwh,
        soc_initial_kwh=p.soc_initial_kwh,
        max_power_kw=p.max_power_kw,
        eta_charge=p.eta_charge,
        eta_discharge=p.eta_discharge,
        lexicographic_abs_tolerance_yuan=p.lexicographic_abs_tolerance_yuan,
        lexicographic_rel_tolerance=p.lexicographic_rel_tolerance,
        emergency_multiplier=p.emergency_multiplier,
        up_adjust_multiplier=p.up_adjust_multiplier,
        down_cancel_refund_multiplier=p.down_cancel_refund_multiplier,
        tolerance=p.tolerance,
    )
    plan = _dispatch_milp(
        np.maximum(net_hat, 0.0), np.maximum(-net_hat, 0.0),
        price, p.soc_initial_kwh, p.soc_initial_kwh, hourly_params,
    )
    actual_net = _hourly_sum(data.load_kwh[day_i] - data.pv_kwh[day_i])
    balance = plan.grid + plan.discharge - plan.charge - actual_net
    emergency = np.maximum(-balance, 0.0)
    return float(np.sum(price * plan.grid + p.emergency_multiplier * price * emergency))


def _empirical_cvar(costs: np.ndarray, level: float) -> float:
    count = max(1, int(np.ceil((1.0 - level) * len(costs))))
    return float(np.mean(np.sort(costs)[-count:]))


def select_dynamic_quantile(
    history: dict[float, list[float]],
    previous_alpha: float,
    cfg: DynamicQuantileParameters,
) -> tuple[float, dict[float, float]]:
    """按近期平均成本、尾部成本和切换惩罚选择下一日净负荷分位数。"""
    scores: dict[float, float] = {}
    for alpha in cfg.candidates:
        costs = np.asarray(history[alpha][-cfg.validation_days:], dtype=float)
        if len(costs) < cfg.validation_days:
            scores[alpha] = np.inf
            continue
        mean_cost = float(np.mean(costs))
        risk_cost = _empirical_cvar(costs, cfg.cvar_level)
        stability = cfg.stability_weight * mean_cost * abs(alpha - previous_alpha) / 0.05
        scores[alpha] = mean_cost + cfg.risk_weight * risk_cost + stability
    finite = {alpha: score for alpha, score in scores.items() if np.isfinite(score)}
    return (min(finite, key=finite.get) if finite else cfg.default_alpha), scores


def forecast_load_median(data: InputData, day_i: int, p: Parameters, issue_t: int) -> np.ndarray:
    """问题3负荷预测：相似日中位数 + 当天已观测误差的指数衰减校正。"""
    candidates = causal_candidates(day_i, data.dates, p.lookback_days)
    base = data.fallback_load_kwh.copy() if len(candidates) == 0 else np.median(data.load_kwh[candidates], axis=0)
    if issue_t > 0:
        window_start = max(0, issue_t - 12)
        # 用 np.mean 而非 sum()/n，避免空窗口和手写除法造成边界错误。
        recent_bias = float(np.mean(data.load_kwh[day_i, window_start:issue_t] - base[window_start:issue_t]))
        horizon = np.arange(144) - issue_t
        base = base + recent_bias * np.exp(-np.maximum(horizon, 0) / 36.0)
    return np.maximum(base, 0.0)


def official_pv_forecast(data: InputData, day_i: int, issue_t: int) -> np.ndarray:
    """把最近一次官方整点预报线性插值到 10 分钟，并用当前实测误差做因果偏差校正。"""
    issue_hour = 0.0 if issue_t == 0 else _time_start_minutes(data.time_labels[issue_t]) / 60.0
    release_hour = int(issue_hour // 6) * 6
    release_index = min(release_hour // 6, 3)
    hourly_kw = data.official_pv_kw[day_i, release_index]
    release_t = _issue_index(data.time_labels, release_hour)
    anchor_kw = 0.0 if release_t == 0 else data.pv_kwh[day_i, release_t - 1] / (1.0 / 6.0)
    point_hours = np.arange(25, dtype=float)
    point_values = np.r_[anchor_kw, hourly_kw]
    start_hours = np.array([_time_start_minutes(x) for x in data.time_labels], dtype=float) / 60.0
    predicted_kw = np.interp(start_hours - release_hour, point_hours, point_values)

    if issue_t > release_t:
        observed_kw = data.pv_kwh[day_i, issue_t - 1] / (1.0 / 6.0)
        predicted_now = float(np.interp(issue_hour - release_hour, point_hours, point_values))
        bias_kw = observed_kw - predicted_now
        future_steps = np.maximum(np.arange(144) - issue_t, 0)
        predicted_kw += bias_kw * np.exp(-future_steps / 24.0)
    return np.maximum(predicted_kw, 0.0) * (1.0 / 6.0)


def _dispatch_milp(
    load_hat: np.ndarray,
    pv_hat: np.ndarray,
    price: np.ndarray,
    initial_soc: float,
    terminal_soc: float | None,
    p: Parameters,
    base_grid: np.ndarray | None = None,
) -> OptimizationResult:
    """两层词典序 MILP：先最小化题目费用，再在最优容差内最小化吞吐量。"""
    n = len(load_hat)
    if not (len(pv_hat) == len(price) == n and n > 0):
        raise ValueError("优化输入长度必须一致且非空")

    adjusted = base_grid is not None
    # 日前变量：[grid, charge, discharge, spill, soc, u_charge, u_discharge]
    # 调整变量：[up, down, charge, discharge, spill, soc, u_charge, u_discharge]
    blocks = 8 if adjusted else 7
    m = blocks * n
    idx = {
        "first": slice(0, n),
        "second": slice(n, 2 * n),
        "charge": slice((2 if adjusted else 1) * n, (3 if adjusted else 2) * n),
        "discharge": slice((3 if adjusted else 2) * n, (4 if adjusted else 3) * n),
        "spill": slice((4 if adjusted else 3) * n, (5 if adjusted else 4) * n),
        "soc": slice((5 if adjusted else 4) * n, (6 if adjusted else 5) * n),
        "u_charge": slice((6 if adjusted else 5) * n, (7 if adjusted else 6) * n),
        "u_discharge": slice((7 if adjusted else 6) * n, blocks * n),
    }

    primary_cost = np.zeros(m)
    lower = np.zeros(m)
    upper = np.full(m, np.inf)
    max_e = p.max_interval_energy_kwh
    grid_cap = np.maximum(load_hat + max_e, max_e)
    if adjusted:
        assert base_grid is not None
        primary_cost[idx["first"]] = p.up_adjust_multiplier * price
        primary_cost[idx["second"]] = -p.down_cancel_refund_multiplier * price
        for t in range(n):
            upper[idx["first"].start + t] = float(grid_cap[t])
            upper[idx["second"].start + t] = float(max(base_grid[t], 0.0))
    else:
        primary_cost[idx["first"]] = price
        for t in range(n):
            upper[idx["first"].start + t] = float(grid_cap[t])

    for t in range(n):
        upper[idx["charge"].start + t] = max_e
        upper[idx["discharge"].start + t] = max_e
        # 剩余量最多等于预测光伏，禁止通过“丢弃已购电”虚假获利。
        upper[idx["spill"].start + t] = float(max(pv_hat[t], 0.0))
        lower[idx["soc"].start + t] = p.soc_min_kwh
        upper[idx["soc"].start + t] = p.soc_max_kwh
        upper[idx["u_charge"].start + t] = 1.0
        upper[idx["u_discharge"].start + t] = 1.0

    eq_rows = 2 * n + (1 if terminal_soc is not None else 0)
    a_eq = lil_matrix((eq_rows, m), dtype=float)
    b_eq = np.zeros(eq_rows)
    for t in range(n):
        # grid + discharge - charge - spill = load - pv
        a_eq[t, idx["first"].start + t] = 1.0
        if adjusted:
            a_eq[t, idx["second"].start + t] = -1.0
            b_eq[t] = load_hat[t] - pv_hat[t] - float(base_grid[t])
        else:
            b_eq[t] = load_hat[t] - pv_hat[t]
        a_eq[t, idx["discharge"].start + t] = 1.0
        a_eq[t, idx["charge"].start + t] = -1.0
        a_eq[t, idx["spill"].start + t] = -1.0

        row = n + t
        a_eq[row, idx["soc"].start + t] = 1.0
        if t > 0:
            a_eq[row, idx["soc"].start + t - 1] = -1.0
            b_eq[row] = 0.0
        else:
            b_eq[row] = initial_soc
        a_eq[row, idx["charge"].start + t] = -p.eta_charge
        a_eq[row, idx["discharge"].start + t] = 1.0 / p.eta_discharge

    if terminal_soc is not None:
        a_eq[2 * n, idx["soc"].stop - 1] = 1.0
        b_eq[2 * n] = terminal_soc

    # 充放电量与状态二元变量联动，并严格禁止同一时段同时充放电。
    a_ub = lil_matrix((3 * n, m), dtype=float)
    ub = np.zeros(3 * n)
    for t in range(n):
        a_ub[t, idx["charge"].start + t] = 1.0
        a_ub[t, idx["u_charge"].start + t] = -max_e
        a_ub[n + t, idx["discharge"].start + t] = 1.0
        a_ub[n + t, idx["u_discharge"].start + t] = -max_e
        a_ub[2 * n + t, idx["u_charge"].start + t] = 1.0
        a_ub[2 * n + t, idx["u_discharge"].start + t] = 1.0
        ub[2 * n + t] = 1.0

    integrality = np.zeros(m, dtype=int)
    integrality[idx["u_charge"]] = 1
    integrality[idx["u_discharge"]] = 1
    bounds = Bounds(lower, upper)
    base_constraints = [
        LinearConstraint(a_eq.tocsr(), b_eq, b_eq),
        LinearConstraint(a_ub.tocsr(), -np.inf, ub),
    ]
    options = {"mip_rel_gap": 1e-9}
    stage1 = milp(primary_cost, integrality=integrality, bounds=bounds, constraints=base_constraints, options=options)
    if not stage1.success:
        raise RuntimeError(f"第一层 HiGHS MILP 求解失败：{stage1.message}")

    cost_tolerance = max(
        p.lexicographic_abs_tolerance_yuan,
        abs(float(stage1.fun)) * p.lexicographic_rel_tolerance,
    )
    throughput_cost = np.zeros(m)
    throughput_cost[idx["charge"]] = 1.0
    throughput_cost[idx["discharge"]] = 1.0
    lex_constraint = LinearConstraint(primary_cost.reshape(1, -1), -np.inf, float(stage1.fun) + cost_tolerance)
    stage2 = milp(
        throughput_cost,
        integrality=integrality,
        bounds=bounds,
        constraints=[*base_constraints, lex_constraint],
        options=options,
    )
    if not stage2.success:
        raise RuntimeError(f"第二层 HiGHS MILP 求解失败：{stage2.message}")

    x = stage2.x
    if adjusted:
        grid = np.maximum(base_grid + x[idx["first"]] - x[idx["second"]], 0.0)
    else:
        grid = x[idx["first"]]
    return OptimizationResult(
        grid=grid,
        charge=x[idx["charge"]],
        discharge=x[idx["discharge"]],
        spill=x[idx["spill"]],
        soc_end=x[idx["soc"]],
        objective=float(primary_cost @ x),
    )


def execute_segment(
    data: InputData,
    day_i: int,
    start: int,
    stop: int,
    grid: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    soc_start: float,
    p: Parameters,
) -> dict[str, np.ndarray | float]:
    """按计划动作执行一段；实际缺口由紧急购电补足，剩余电量单独记录。"""
    n = stop - start
    soc = np.zeros(n)
    emergency = np.zeros(n)
    spill = np.zeros(n)
    current = soc_start
    for j, t in enumerate(range(start, stop)):
        c = float(charge[j])
        d = float(discharge[j])
        current = current + p.eta_charge * c - d / p.eta_discharge
        if not (p.soc_min_kwh - p.tolerance <= current <= p.soc_max_kwh + p.tolerance):
            raise AssertionError(f"{data.dates[day_i].date()} 时段 {t} SOC 越界：{current}")
        balance = grid[j] + data.pv_kwh[day_i, t] + d - data.load_kwh[day_i, t] - c
        emergency[j] = max(-balance, 0.0)
        spill[j] = max(balance, 0.0)
        soc[j] = current
    return {"soc": soc, "emergency": emergency, "spill": spill, "end_soc": current}


def _daily_records(
    data: InputData,
    day_i: int,
    price: np.ndarray,
    q0: np.ndarray,
    q_effective: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    soc: np.ndarray,
    emergency: np.ndarray,
    spill: np.ndarray,
    start_soc: float,
    p: Parameters,
    strategy: str,
) -> pd.DataFrame:
    up = np.maximum(q_effective - q0, 0.0)
    down = np.maximum(q0 - q_effective, 0.0)
    plan_cost = price * q0
    adjustment_cost = price * (p.up_adjust_multiplier * up - p.down_cancel_refund_multiplier * down)
    emergency_cost = p.emergency_multiplier * price * emergency
    return pd.DataFrame({
        "策略": strategy,
        "日期": data.dates[day_i],
        "时段序号": np.arange(144),
        "时间标签": data.time_labels,
        "电价_元每kWh": price,
        "实际负荷_kWh": data.load_kwh[day_i],
        "实际光伏_kWh": data.pv_kwh[day_i],
        "计划购电量_kWh": q0,
        "调整后购电量_kWh": q_effective,
        "上调量_kWh": up,
        "下调量_kWh": down,
        "充电量_kWh": charge,
        "放电量_kWh": discharge,
        "期末SOC_kWh": soc,
        "紧急购电量_kWh": emergency,
        "剩余电量_kWh": spill,
        "计划购电费_元": plan_cost,
        "调整成本_元": adjustment_cost,
        "紧急购电费_元": emergency_cost,
        "总成本_元": plan_cost + adjustment_cost + emergency_cost,
        "期初SOC_kWh": np.r_[start_soc, soc[:-1]],
    })


def solve_day_ahead(
    data: InputData,
    p: Parameters,
    price_mode: str,
    stop_day: int,
) -> pd.DataFrame:
    """问题2/4-2：每天滚动选取净负荷分位数，再制定全天计划。"""
    strategy = "问题2_固定价" if price_mode == "fixed" else "问题4-2_波动价"
    cfg = DynamicQuantileParameters()
    history = {alpha: [] for alpha in cfg.candidates}
    selected_alpha = cfg.default_alpha
    rows: list[pd.DataFrame] = []
    soc = p.soc_initial_kwh
    for day_i in range(stop_day):
        price = data.fixed_price if price_mode == "fixed" else data.variable_price[day_i]
        selected_alpha, scores = select_dynamic_quantile(history, selected_alpha, cfg)
        net_hat = dynamic_net_forecast(data, day_i, selected_alpha, cfg)
        load_hat = np.maximum(net_hat, 0.0)
        pv_hat = np.maximum(-net_hat, 0.0)
        terminal_soc = p.soc_initial_kwh if day_i == stop_day - 1 else None
        plan = _dispatch_milp(load_hat, pv_hat, price, soc, terminal_soc, p)
        executed = execute_segment(data, day_i, 0, 144, plan.grid, plan.charge, plan.discharge, soc, p)
        frame = _daily_records(
            data, day_i, price, plan.grid, plan.grid, plan.charge, plan.discharge,
            executed["soc"], executed["emergency"], executed["spill"], soc, p, strategy,
        )
        frame["选择净负荷分位数"] = selected_alpha
        frame["分位数选择得分"] = scores.get(selected_alpha, np.nan)
        rows.append(frame)
        soc = float(executed["end_soc"])
        # 当天执行结束后才加入当天反事实成本，确保次日选参严格因果。
        for alpha in cfg.candidates:
            history[alpha].append(
                dynamic_validation_cost(data, day_i, alpha, cfg, p, price_mode)
            )
    return pd.concat(rows, ignore_index=True)


def solve_rolling(
    data: InputData,
    p: Parameters,
    price_mode: str,
    stop_day: int,
    issue_hours: tuple[int, ...] = (0, 6, 12, 18),
    strategy_suffix: str = "",
) -> pd.DataFrame:
    """问题3/4-3：在每个信息时点冻结历史，只重算尚未执行的时段。"""
    base_name = "问题3_固定价" if price_mode == "fixed" else "问题4-3_波动价"
    strategy = base_name + strategy_suffix
    rows: list[pd.DataFrame] = []
    soc = p.soc_initial_kwh
    issue_ts = tuple(_issue_index(data.time_labels, h) for h in issue_hours)
    if issue_ts[0] != 0 or tuple(sorted(set(issue_ts))) != issue_ts:
        raise ValueError("issue_hours 必须从 0 开始且严格递增")

    for day_i in range(stop_day):
        price = data.fixed_price if price_mode == "fixed" else data.variable_price[day_i]
        load0 = forecast_load_median(data, day_i, p, 0)
        pv0 = official_pv_forecast(data, day_i, 0)
        terminal_soc = p.soc_initial_kwh if day_i == stop_day - 1 else None
        plan0 = _dispatch_milp(load0, pv0, price, soc, terminal_soc, p)
        q0 = plan0.grid.copy()
        q_final = q0.copy()
        charge_final = plan0.charge.copy()
        discharge_final = plan0.discharge.copy()
        soc_final = np.zeros(144)
        emergency_final = np.zeros(144)
        spill_final = np.zeros(144)
        day_start_soc = soc

        for k, start in enumerate(issue_ts):
            stop = issue_ts[k + 1] if k + 1 < len(issue_ts) else 144
            if start == 0:
                current_plan = plan0
            else:
                load_hat = forecast_load_median(data, day_i, p, start)[start:]
                pv_hat = official_pv_forecast(data, day_i, start)[start:]
                current_plan = _dispatch_milp(
                    load_hat, pv_hat, price[start:], soc, terminal_soc, p, base_grid=q0[start:]
                )
                q_final[start:] = current_plan.grid
                charge_final[start:] = current_plan.charge
                discharge_final[start:] = current_plan.discharge

            local_stop = stop - start
            executed = execute_segment(
                data, day_i, start, stop,
                current_plan.grid[:local_stop], current_plan.charge[:local_stop],
                current_plan.discharge[:local_stop], soc, p,
            )
            q_final[start:stop] = current_plan.grid[:local_stop]
            charge_final[start:stop] = current_plan.charge[:local_stop]
            discharge_final[start:stop] = current_plan.discharge[:local_stop]
            soc_final[start:stop] = executed["soc"]
            emergency_final[start:stop] = executed["emergency"]
            spill_final[start:stop] = executed["spill"]
            soc = float(executed["end_soc"])

        rows.append(_daily_records(
            data, day_i, price, q0, q_final, charge_final, discharge_final,
            soc_final, emergency_final, spill_final, day_start_soc, p, strategy,
        ))
    return pd.concat(rows, ignore_index=True)


def validate_schedule(frame: pd.DataFrame, p: Parameters) -> dict[str, float | int | bool]:
    """结果输出前独立回代，不以优化器 success 代替物理校验。"""
    ordered = frame.sort_values(["日期", "时段序号"]).reset_index(drop=True)
    balance = (
        frame["调整后购电量_kWh"] + frame["实际光伏_kWh"] + frame["放电量_kWh"]
        + frame["紧急购电量_kWh"] - frame["实际负荷_kWh"]
        - frame["充电量_kWh"] - frame["剩余电量_kWh"]
    )
    soc_residual = (
        frame["期末SOC_kWh"] - frame["期初SOC_kWh"]
        - p.eta_charge * frame["充电量_kWh"] + frame["放电量_kWh"] / p.eta_discharge
    )
    simultaneous = (frame["充电量_kWh"] > 1e-5) & (frame["放电量_kWh"] > 1e-5)
    cost_rebuilt = frame[["计划购电费_元", "调整成本_元", "紧急购电费_元"]].sum(axis=1)
    daily_soc = ordered.groupby("日期", sort=True).agg(
        start=("期初SOC_kWh", "first"), end=("期末SOC_kWh", "last")
    )
    cross_day_residual = np.abs(
        daily_soc["start"].iloc[1:].to_numpy() - daily_soc["end"].iloc[:-1].to_numpy()
    )
    checks = {
        "最大供需平衡残差_kWh": float(np.abs(balance).max()),
        "最大SOC递推残差_kWh": float(np.abs(soc_residual).max()),
        "最小SOC_kWh": float(frame["期末SOC_kWh"].min()),
        "最大SOC_kWh": float(frame["期末SOC_kWh"].max()),
        "同时充放电时段数": int(simultaneous.sum()),
        "最大成本对账误差_元": float(np.abs(cost_rebuilt - frame["总成本_元"]).max()),
        "最大跨日SOC衔接残差_kWh": float(cross_day_residual.max()) if len(cross_day_residual) else 0.0,
        "计算区间初始SOC_kWh": float(daily_soc["start"].iloc[0]),
        "计算区间最终SOC_kWh": float(daily_soc["end"].iloc[-1]),
        "非末日首末SOC不同天数": int(
            ((daily_soc["end"].iloc[:-1] - daily_soc["start"].iloc[:-1]).abs() > 1e-6).sum()
        ),
        "日期数": int(frame["日期"].nunique()),
    }
    if checks["最大供需平衡残差_kWh"] > 1e-5 or checks["最大SOC递推残差_kWh"] > 1e-5:
        raise AssertionError(f"平衡或 SOC 回代失败：{checks}")
    if checks["最小SOC_kWh"] < p.soc_min_kwh - 1e-5 or checks["最大SOC_kWh"] > p.soc_max_kwh + 1e-5:
        raise AssertionError(f"SOC 边界失败：{checks}")
    if checks["同时充放电时段数"] != 0:
        raise AssertionError(f"出现同时充放电：{checks}")
    if checks["最大成本对账误差_元"] > 1e-6:
        raise AssertionError(f"成本对账失败：{checks}")
    if checks["最大跨日SOC衔接残差_kWh"] > 1e-6:
        raise AssertionError(f"跨日 SOC 衔接失败：{checks}")
    if abs(checks["计算区间初始SOC_kWh"] - p.soc_initial_kwh) > 1e-5:
        raise AssertionError(f"计算区间初始 SOC 错误：{checks}")
    if abs(checks["计算区间最终SOC_kWh"] - p.soc_initial_kwh) > 1e-5:
        raise AssertionError(f"计算区间最终 SOC 错误：{checks}")
    return checks


def summarize(frame: pd.DataFrame) -> pd.DataFrame:
    daily = frame.groupby(["策略", "日期"], as_index=False).agg(
        计划购电量_kWh=("计划购电量_kWh", "sum"),
        调整后购电量_kWh=("调整后购电量_kWh", "sum"),
        上调量_kWh=("上调量_kWh", "sum"),
        下调量_kWh=("下调量_kWh", "sum"),
        紧急购电量_kWh=("紧急购电量_kWh", "sum"),
        剩余电量_kWh=("剩余电量_kWh", "sum"),
        计划购电费_元=("计划购电费_元", "sum"),
        调整成本_元=("调整成本_元", "sum"),
        紧急购电费_元=("紧急购电费_元", "sum"),
        总成本_元=("总成本_元", "sum"),
    )
    return daily


def interval_label(t: int) -> str:
    def fmt(minutes: int) -> str:
        day = minutes // 1440
        minute = minutes % 1440
        return f"{minute // 60}:{minute % 60:02d}" + ("+1" if day else "")
    return f"{fmt((t + 1) * 10)}-{fmt((t + 2) * 10)}"


def emergency_groups(day_frame: pd.DataFrame, tol: float = 1e-6) -> list[tuple[str, float]]:
    values = day_frame["紧急购电量_kWh"].to_numpy(float)
    groups: list[tuple[str, float]] = []
    start: int | None = None
    for t, value in enumerate(np.r_[values, 0.0]):
        if value > tol and start is None:
            start = t
        if value <= tol and start is not None:
            groups.append((f"{interval_label(start).split('-')[0]}-{interval_label(t - 1).split('-')[1]}", float(values[start:t].sum())))
            start = None
    return groups


def _capture_row_style(ws, source_row: int, max_col: int) -> list:
    """删除模板占位行前保存一份真实数据行样式。"""
    return [copy(ws.cell(source_row, col)._style) for col in range(1, max_col + 1)]


def _apply_row_style(ws, styles: list, target_row: int) -> None:
    for col, style in enumerate(styles, start=1):
        ws.cell(target_row, col)._style = copy(style)


def write_template(frame: pd.DataFrame, template_name: str, output_name: str, adjusted: bool) -> Path:
    """保持官方工作表名称和表头，扩展模板占位行并写入全部日期结果。"""
    output_path = RESULTS_DIR / output_name
    shutil.copy2(TEMPLATE_DIR / template_name, output_path)
    wb = load_workbook(output_path)
    all_rows = frame.copy()
    output = all_rows[all_rows["日期"] >= pd.Timestamp("2025-02-01")].copy()
    dates = sorted(output["日期"].unique())

    def write_purchase(sheet_name: str, value_col: str, cost_mode: str) -> None:
        ws = wb[sheet_name]
        for row_i, date in enumerate(dates, start=2):
            day = output[output["日期"] == date].sort_values("时段序号")
            ws.cell(row_i, 1, pd.Timestamp(date).to_pydatetime())
            for col_i, value in enumerate(day[value_col].to_numpy(float), start=2):
                ws.cell(row_i, col_i, round(float(value), 6))
            ws.cell(row_i, 146, round(float(day[value_col].sum()), 6))
            if cost_mode == "plan":
                cost = day["计划购电费_元"].sum()
            else:
                cost = day["总成本_元"].sum()
            ws.cell(row_i, 147, round(float(cost), 6))
        ws.freeze_panes = "B2"

    write_purchase("计划购电量", "计划购电量_kWh", "plan")
    if adjusted:
        write_purchase("调整购电量", "调整后购电量_kWh", "total")

    ws = wb["充放电量"]
    storage_row_style = _capture_row_style(ws, 2, 6)
    ws.delete_rows(2, max(ws.max_row - 1, 1))
    row = 2
    for date in dates:
        day = output[output["日期"] == date].sort_values("时段序号")
        previous_day = all_rows[all_rows["日期"] == pd.Timestamp(date) - pd.Timedelta(days=1)].sort_values("时段序号")
        previous_midnight = previous_day[
            previous_day["时间标签"].map(_time_start_minutes) == 1440
        ]
        current_midnight = day[day["时间标签"].map(_time_start_minutes) == 1440]
        if len(previous_midnight) != 1 or len(current_midnight) != 1:
            raise ValueError(f"{pd.Timestamp(date).date()} 无法定位自然日 0:00/24:00 的 SOC")
        day_start = float(previous_midnight["期初SOC_kWh"].iloc[0])
        day_end = float(current_midnight["期初SOC_kWh"].iloc[0])
        for block in range(6):
            _apply_row_style(ws, storage_row_style, row)
            clock_minutes = day["时间标签"].map(_time_start_minutes) % 1440
            if block == 0:
                # 当前日期 0:00-0:10 位于附件上一日期的末列；其余 23 段来自当前日期。
                segment = pd.concat([
                    previous_midnight,
                    day[(clock_minutes >= 10) & (clock_minutes < 240)],
                ])
            else:
                segment = day[(clock_minutes >= block * 240) & (clock_minutes < (block + 1) * 240)]
            if len(segment) != 24:
                raise ValueError(f"{pd.Timestamp(date).date()} 的 {4*block}:00-{4*(block+1)}:00 不是24个时段")
            ws.cell(row, 1, pd.Timestamp(date).to_pydatetime() if block == 0 else None)
            ws.cell(row, 2, f"{4 * block}:00-{4 * (block + 1)}:00")
            ws.cell(row, 3, round(float(segment["充电量_kWh"].sum()), 6))
            ws.cell(row, 4, round(float(segment["放电量_kWh"].sum()), 6))
            if block == 0:
                ws.cell(row, 5, "0:00")
                ws.cell(row, 6, round(day_start, 6))
            elif block == 1:
                ws.cell(row, 5, "24:00")
                ws.cell(row, 6, round(day_end, 6))
            row += 1
    ws.freeze_panes = "A2"

    ws = wb["紧急购电量"]
    emergency_row_style = _capture_row_style(ws, 2, 3)
    ws.delete_rows(2, max(ws.max_row - 1, 1))
    row = 2
    for date in dates:
        day = output[output["日期"] == date].sort_values("时段序号")
        groups = emergency_groups(day)
        if not groups:
            groups = [("无", 0.0)]
        for j, (label, value) in enumerate(groups):
            _apply_row_style(ws, emergency_row_style, row)
            ws.cell(row, 1, pd.Timestamp(date).to_pydatetime() if j == 0 else None)
            ws.cell(row, 2, label)
            ws.cell(row, 3, round(value, 6))
            row += 1
    ws.freeze_panes = "A2"

    for ws in wb.worksheets:
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
    wb.save(output_path)
    return output_path


def add_caption(fig, text: str) -> None:
    fig.tight_layout(rect=(0, 0.09, 1, 1))
    fig.text(0.5, 0.018, text, ha="center", va="bottom", fontsize=9)


def setup_plot_style() -> None:
    candidates = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"]
    available = {x.name for x in __import__("matplotlib").font_manager.fontManager.ttflist}
    font = next((x for x in candidates if x in available), "DejaVu Sans")
    plt.rcParams.update({
        "font.family": "sans-serif", "font.sans-serif": [font], "axes.unicode_minus": False,
        "figure.dpi": 130, "savefig.dpi": 300, "axes.spines.top": False,
        "axes.spines.right": False, "axes.grid": True, "grid.alpha": 0.18,
    })


def make_figures(all_frame: pd.DataFrame, daily: pd.DataFrame, voi_daily: pd.DataFrame | None) -> list[Path]:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    setup_plot_style()
    colors = ["#0072B2", "#009E73", "#D55E00", "#CC79A7"]
    paths: list[Path] = []
    official = daily[daily["日期"] >= pd.Timestamp("2025-02-01")].copy()

    total = official.groupby("策略", as_index=False)["总成本_元"].sum().sort_values("总成本_元")
    fig, ax = plt.subplots(figsize=(8.2, 5.0))
    bars = ax.bar(total["策略"], total["总成本_元"] / 1e6, color=colors[:len(total)], label="全年总成本")
    ax.bar_label(bars, fmt="%.2f", padding=3)
    ax.set(title="问题2至问题4全年购电总成本", xlabel="策略", ylabel="总成本（百万元）")
    ax.legend()
    saving = float(total["总成本_元"].max() - total["总成本_元"].min()) / 1e6
    add_caption(fig, f"关键结论：最高与最低成本方案相差 {saving:.2f} 百万元。")
    path = FIGURES_DIR / "result_q234_total_cost.png"; fig.savefig(path, bbox_inches="tight"); plt.close(fig); paths.append(path)

    official["月份"] = official["日期"].dt.to_period("M").astype(str)
    monthly = official.groupby(["策略", "月份"], as_index=False)["总成本_元"].sum()
    fig, ax = plt.subplots(figsize=(9.0, 5.2))
    for color, (name, group) in zip(colors, monthly.groupby("策略", sort=False)):
        ax.plot(group["月份"], group["总成本_元"] / 1e6, label=name, color=color, marker="o", markevery=2)
    ax.set(title="月度总成本趋势", xlabel="月份", ylabel="总成本（百万元）")
    ax.tick_params(axis="x", rotation=35); ax.legend(ncol=2)
    peak = monthly.loc[monthly["总成本_元"].idxmax()]
    add_caption(fig, f"关键结论：月度最高成本出现在 {peak['月份']}，对应 {peak['策略']}。")
    path = FIGURES_DIR / "result_q234_monthly_cost.png"; fig.savefig(path, bbox_inches="tight"); plt.close(fig); paths.append(path)

    emergency = official.groupby("策略", as_index=False)["紧急购电量_kWh"].sum()
    fig, ax = plt.subplots(figsize=(8.2, 5.0))
    bars = ax.bar(emergency["策略"], emergency["紧急购电量_kWh"] / 1e3, color=colors[:len(emergency)], label="紧急购电量")
    ax.bar_label(bars, fmt="%.1f", padding=3)
    ax.set(title="各策略紧急购电量", xlabel="策略", ylabel="紧急购电量（MWh）")
    ax.legend()
    best = emergency.loc[emergency["紧急购电量_kWh"].idxmin()]
    add_caption(fig, f"关键结论：{best['策略']} 的紧急购电量最低，为 {best['紧急购电量_kWh']/1e3:.1f} MWh。")
    path = FIGURES_DIR / "result_q234_emergency_energy.png"; fig.savefig(path, bbox_inches="tight"); plt.close(fig); paths.append(path)

    example_date = pd.Timestamp("2025-09-23")
    example = all_frame[(all_frame["策略"] == "问题4-3_波动价") & (all_frame["日期"] == example_date)].sort_values("时段序号")
    if len(example) == 144:
        x = np.array([_time_start_minutes(v) for v in example["时间标签"]], dtype=float) / 60.0
        fig, axes = plt.subplots(3, 1, figsize=(9.0, 7.5), sharex=True)
        axes[0].plot(x, example["电价_元每kWh"], label="波动电价", color=colors[2])
        axes[0].set(title="典型日价格", ylabel="元/kWh"); axes[0].legend()
        axes[1].plot(x, example["计划购电量_kWh"], label="0:00计划", color=colors[0])
        axes[1].plot(x, example["调整后购电量_kWh"], label="最终调整", color=colors[1], linestyle="--")
        axes[1].fill_between(x, 0, example["紧急购电量_kWh"], label="紧急购电", color=colors[2], alpha=0.35)
        axes[1].set(title="购电计划与实际补救", ylabel="kWh/时段"); axes[1].legend(ncol=3)
        axes[2].plot(x, example["期末SOC_kWh"], label="SOC", color=colors[3])
        axes[2].axhline(p_global.soc_min_kwh, color="#666666", linestyle=":", label="安全边界")
        axes[2].axhline(p_global.soc_max_kwh, color="#666666", linestyle=":")
        axes[2].set(title="储能状态", xlabel="时刻（h）", ylabel="kWh"); axes[2].legend()
        adjusted_energy = float(example["上调量_kWh"].sum() + example["下调量_kWh"].sum())
        add_caption(fig, f"关键结论：{example_date.date()} 最终调整总量为 {adjusted_energy:.1f} kWh，SOC 全程满足安全边界。")
        path = FIGURES_DIR / "process_q43_typical_day.png"; fig.savefig(path, bbox_inches="tight"); plt.close(fig); paths.append(path)

    if voi_daily is not None and not voi_daily.empty:
        base_cost = float(official[official["策略"] == "问题3_固定价"]["总成本_元"].sum())
        dense_cost = float(voi_daily[voi_daily["日期"] >= pd.Timestamp("2025-02-01")]["总成本_元"].sum())
        values = np.array([base_cost, dense_cost]) / 1e6
        fig, ax = plt.subplots(figsize=(7.5, 4.8))
        bars = ax.bar(["四时点", "八时点代理"], values, color=[colors[0], colors[1]], label="总成本")
        ax.bar_label(bars, fmt="%.3f", padding=3)
        ax.set(title="增加预测/更新时点的信息价值上界", xlabel="更新方案", ylabel="总成本（百万元）")
        ax.legend()
        voi = base_cost - dense_cost
        add_caption(fig, f"关键结论：增加 3/9/15/21 点因果更新的毛信息价值为 {voi/1e4:.2f} 万元；实际部署还应扣除预报与通信成本。")
        path = FIGURES_DIR / "result_q3_information_value.png"; fig.savefig(path, bbox_inches="tight"); plt.close(fig); paths.append(path)
    return paths


def write_notes(
    p: Parameters,
    frame: pd.DataFrame,
    summaries: pd.DataFrame,
    checks: dict,
    voi_value: float | None,
) -> Path:
    totals = summaries[summaries["日期"] >= pd.Timestamp("2025-02-01")].groupby("策略").sum(numeric_only=True)
    dynamic_cfg = DynamicQuantileParameters()
    evaluation = frame[frame["日期"] >= pd.Timestamp("2025-02-01")]
    quantile_stats: dict[str, tuple[float, float, int]] = {}
    for name in ("问题2_固定价", "问题4-2_波动价"):
        selected = evaluation[evaluation["策略"] == name].groupby("日期")["选择净负荷分位数"].first()
        quantile_stats[name] = (
            float(selected.mean()),
            float(selected.median()),
            int(selected.ne(selected.shift()).sum() - 1),
        )
    lines = [
        "# 问题2至问题4代码求解说明",
        "",
        "## 依赖与运行",
        "",
        "```powershell",
        "python -m pip install numpy pandas scipy matplotlib openpyxl pillow",
        "python q234_solve.py --smoke  # 先做快速检查",
        "python q234_solve.py          # 再生成全年正式结果",
        "```",
        "",
        "## 统一求解步骤",
        "",
        "1. 数据输入：读取附件1至附件4，逐一核对日期、144个时段和非负数值。注意：kW 必须乘 `1/6 h` 才能得到 kWh。",
        "2. 参数初始化：设置 SOC 范围、功率上限、效率、五倍紧急电价和动态净负荷分位数候选集。注意：效率必须在 `(0,1]`，候选分位数必须在 `[0,1]`。",
        "3. 模型调用：问题2按近期样本外运行成本滚动选择净负荷分位数并求解日前混合整数规划；问题3在0/6/12/18点重算未来时段；问题4替换为附件4波动电价。注意：历史时段不得重算。",
        "4. 结果输出：回代供需平衡、SOC递推、边界、互斥和成本，再填写官方模板并生成图表。注意：不得只根据求解器 success 判断结果正确。",
        "",
        "## 关键口径",
        "",
        f"- 单程充、放电效率均为 {p.eta_charge:.2f}，往返效率为 {p.eta_charge*p.eta_discharge:.2%}。",
        f"- 问题2与问题4-2直接预测净负荷：训练窗口 {dynamic_cfg.training_days} 天、相似日 {dynamic_cfg.similar_days} 个、验证窗口 {dynamic_cfg.validation_days} 天。",
        f"- 候选分位数为 {dynamic_cfg.candidates[0]:.2f}—{dynamic_cfg.candidates[-1]:.2f}、步长 {dynamic_cfg.candidates[1]-dynamic_cfg.candidates[0]:.2f}；条件风险价值置信水平 {dynamic_cfg.cvar_level:.2f}、风险权重 {dynamic_cfg.risk_weight:.2f}、稳定权重 {dynamic_cfg.stability_weight:.2f}、初始分位数 {dynamic_cfg.default_alpha:.2f}。",
        "- 负荷与光伏按同一历史日配对后直接构造净负荷，避免分别取边际分位数造成双重保守；候选参数按已实现历史日的运行成本而非单一预测误差选择。",
        "- 动态净负荷分位数只覆盖问题2固定价和问题4-2波动价；问题3和问题4-3使用附件3的0/6/12/18时点光伏预报滚动求解。",
        "- 第三问下调后不再支付被取消电量的原价，但支付其50%违约费；因此总成本为 `原计划费 + 1.5p×上调量 - 0.5p×下调量 + 5p×紧急购电量`。",
        "- 问题4假定当日144点电价在0:00已知。若赛题解释为实时才可见，应替换为只使用历史数据的价格预测。",
        "- 逐时段二元状态变量与有限功率上界严格禁止同时充放电。",
        "- 采用两层词典序优化：第一层只最小化题目规定费用，第二层在第一层最优费用的数值容差内最小化充放电总吞吐量。",
        "- 储能状态跨日连续，不再逐日强制首末相等；仅在整个计算区间最后一天回到初始储电量。",
        "- 当前仍按每日可得信息逐日前推求解，不使用未来实际数据；它是因果滚动策略，不是拥有全年未来信息的一次性联合最优。",
        "",
        "## 全年结果（2025-02-01至2025-12-31）",
        "",
        "| 策略 | 总成本（元） | 计划购电量（kWh） | 调整后购电量（kWh） | 紧急购电量（kWh） |",
        "|---|---:|---:|---:|---:|",
    ]
    for name, row in totals.iterrows():
        lines.append(f"| {name} | {row['总成本_元']:.2f} | {row['计划购电量_kWh']:.2f} | {row['调整后购电量_kWh']:.2f} | {row['紧急购电量_kWh']:.2f} |")
    fixed_saving = totals.loc["问题2_固定价", "总成本_元"] - totals.loc["问题3_固定价", "总成本_元"]
    variable_saving = totals.loc["问题4-2_波动价", "总成本_元"] - totals.loc["问题4-3_波动价", "总成本_元"]
    lines += [
        "", "## 动态选参统计与经济比较", "",
        f"- 问题2所选分位数均值 {quantile_stats['问题2_固定价'][0]:.3f}、中位数 {quantile_stats['问题2_固定价'][1]:.2f}、变更 {quantile_stats['问题2_固定价'][2]} 次。",
        f"- 问题4-2所选分位数均值 {quantile_stats['问题4-2_波动价'][0]:.3f}、中位数 {quantile_stats['问题4-2_波动价'][1]:.2f}、变更 {quantile_stats['问题4-2_波动价'][2]} 次。",
        f"- 问题3相对问题2节省 {fixed_saving:.2f} 元，降幅 {fixed_saving/totals.loc['问题2_固定价', '总成本_元']:.3%}。",
        f"- 问题4-3相对问题4-2节省 {variable_saving:.2f} 元，降幅 {variable_saving/totals.loc['问题4-2_波动价', '总成本_元']:.3%}。",
    ]
    lines += ["", "## 校验结果", ""]
    for name, item in checks.items():
        lines.append(f"- {name}：最大供需残差 {item['最大供需平衡残差_kWh']:.3e} kWh，最大 SOC 残差 {item['最大SOC递推残差_kWh']:.3e} kWh，同时充放电 {item['同时充放电时段数']} 个时段。")
    lines += [
        "- 独立五折时间顺序检验中，动态日前净负荷预测的平均选择分位数为 0.808；其点预测 RMSE 不优于周滞后基线，但分位数按运行成本选取，不能据此直接判定经济方案失效。",
        "- 独立正负10%冻结压力检验中，日前分支平均成本增加约6.6%，滚动分支约增加12.0%至12.3%；这是后续滚动补救失效时的压力上界。",
        "", "## 适用条件与限制", "",
        "- 问题四假定每日0时已知当日完整价格路径；若价格只能实时获得，必须接入只使用历史信息的价格预测器后重新求解。",
        "- 题目未提供并网购电功率上限，当前不额外设置该约束；如补充上限，应重新检查可行性。",
        "- 当前按每日可得信息逐日前推，属于因果滚动策略，不是拥有全年未来信息的一次性联合最优。",
        "- 上述五折与压力检验结论来自当前正式模型检验文件；模型、结果或检验口径改变后须同步重跑，不得沿用旧结论。",
    ]
    if voi_value is not None:
        conclusion = "具有正的毛价值" if voi_value > 0 else "未显示正的毛价值"
        decision = (
            "只有当新增预报、通信与调整系统的全年总成本低于该正毛价值时，才建议增加相应时点。"
            if voi_value > 0
            else "由于毛信息价值为负，即使暂不计新增系统成本，也不建议按当前代理方案增加这些时点。"
        )
        lines += [
            "", "## 是否增加预报时点", "",
            f"在3/9/15/21点利用最新实测偏差更新既有官方预报的八时点代理方案，相对四时点方案的毛信息价值为 {voi_value:.2f} 元，{conclusion}。",
            f"该值是低成本因果更新的经济上界/代理值。{decision}",
        ]
    path = ROOT / "问题2至4_求解说明.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main() -> None:
    global p_global
    args = parse_args()
    p_global = Parameters(
        eta_charge=args.eta,
        eta_discharge=args.eta,
    )
    validate_parameters(p_global)
    data = load_inputs(p_global)
    stop_day = 34 if args.smoke else 365
    tic = time.perf_counter()

    q2 = solve_day_ahead(data, p_global, "fixed", stop_day)
    q42 = solve_day_ahead(data, p_global, "variable", stop_day)
    q3 = solve_rolling(data, p_global, "fixed", stop_day)
    q43 = solve_rolling(data, p_global, "variable", stop_day)
    all_frame = pd.concat([q2, q3, q42, q43], ignore_index=True)
    checks = {name: validate_schedule(group, p_global) for name, group in all_frame.groupby("策略")}
    daily = summarize(all_frame)

    voi_daily = None
    voi_value = None
    if not args.smoke:
        dense = solve_rolling(
            data, p_global, "fixed", stop_day,
            issue_hours=(0, 3, 6, 9, 12, 15, 18, 21), strategy_suffix="_八时点代理",
        )
        dense_check = validate_schedule(dense, p_global)
        checks["问题3_固定价_八时点代理"] = dense_check
        voi_daily = summarize(dense)
        base_cost = daily[(daily["策略"] == "问题3_固定价") & (daily["日期"] >= pd.Timestamp("2025-02-01"))]["总成本_元"].sum()
        dense_cost = voi_daily[voi_daily["日期"] >= pd.Timestamp("2025-02-01")]["总成本_元"].sum()
        voi_value = float(base_cost - dense_cost)

    print("[1/4] 数据输入完成：365天×144时段，所有功率已换算为kWh")
    dynamic_cfg = DynamicQuantileParameters()
    print(
        "[2/4] 参数初始化完成："
        f"eta={p_global.eta_charge:.2f}, "
        f"日前净负荷分位数候选={dynamic_cfg.candidates[0]:.2f}—{dynamic_cfg.candidates[-1]:.2f}, "
        f"验证窗={dynamic_cfg.validation_days}天"
    )
    print(f"[3/4] 模型调用完成：{len(all_frame):,}条逐时段结果")

    if args.smoke:
        print(json.dumps(checks, ensure_ascii=False, indent=2))
        print(f"快速检查通过，用时 {time.perf_counter()-tic:.2f} 秒")
        return

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    all_frame.to_csv(RESULTS_DIR / "问题2至4_逐时段完整结果.csv", index=False, encoding="utf-8-sig")
    daily.to_csv(RESULTS_DIR / "问题2至4_逐日汇总.csv", index=False, encoding="utf-8-sig")
    if voi_daily is not None:
        voi_daily.to_csv(RESULTS_DIR / "问题3_新增预报时点_VOI.csv", index=False, encoding="utf-8-sig")

    template_paths = [
        write_template(q2, "result2.xlsx", "result2.xlsx", adjusted=False),
        write_template(q3, "result3.xlsx", "result3.xlsx", adjusted=True),
        write_template(q42, "result4-2.xlsx", "result4-2.xlsx", adjusted=False),
        write_template(q43, "result4-3.xlsx", "result4-3.xlsx", adjusted=True),
    ]
    figures = make_figures(all_frame, daily, voi_daily)
    notes = write_notes(p_global, all_frame, daily, checks, voi_value)
    summary_payload = {
        "parameters": asdict(p_global),
        "dynamic_quantile_parameters": asdict(DynamicQuantileParameters()),
        "dynamic_quantile_coverage": ["问题2_固定价", "问题4-2_波动价"],
        "checks": checks,
        "voi_yuan": voi_value,
        "elapsed_seconds": time.perf_counter() - tic,
        "result_workbooks": [str(x) for x in template_paths],
        "figures": [str(x) for x in figures],
    }
    (RESULTS_DIR / "问题2至4_运行摘要.json").write_text(
        json.dumps(summary_payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    totals = daily[daily["日期"] >= pd.Timestamp("2025-02-01")].groupby("策略")["总成本_元"].sum()
    print("[4/4] 结果输出与校验完成")
    for name, value in totals.items():
        print(f"  {name}: {value:,.2f} 元")
    if voi_value is not None:
        print(f"  八时点代理毛信息价值: {voi_value:,.2f} 元")
    print(f"  结果说明: {notes}")
    print(f"  总用时: {time.perf_counter()-tic:.2f} 秒")


p_global = Parameters()


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError, RuntimeError, AssertionError) as exc:
        print(f"求解失败：{exc}", file=sys.stderr)
        raise SystemExit(1) from exc
