from pathlib import Path
import unittest

import pandas as pd
from openpyxl import load_workbook

import problem1_solve as q1
import q234_solve as q234
import numpy as np


ROOT = Path(__file__).resolve().parent


class TimeAlignmentTests(unittest.TestCase):
    def test_q234_milp_allows_cross_day_energy_transfer(self):
        params = q234.Parameters()
        result = q234._dispatch_milp(
            load_hat=np.array([100.0, 100.0]),
            pv_hat=np.zeros(2),
            price=np.array([1.0, 10.0]),
            initial_soc=6000.0,
            terminal_soc=None,
            p=params,
        )
        self.assertNotAlmostEqual(float(result.soc_end[-1]), 6000.0, places=5)
        self.assertFalse(np.any((result.charge > 1e-6) & (result.discharge > 1e-6)))

    def test_attachment_one_uses_start_time_labels(self):
        frame = q1.read_input(q1.INPUT_XLSX)
        self.assertEqual(frame["时段起始分钟"].tolist(), list(range(10, 1450, 10)))
        self.assertEqual(q1.interval_label(0), "0:10-0:20")
        self.assertEqual(q1.interval_label(143), "0:00+1-0:10+1")

    def test_rolling_updates_start_at_the_named_hour(self):
        data = q234.load_inputs(q234.Parameters())
        self.assertEqual(
            [q234._issue_index(data.time_labels, hour) for hour in (0, 6, 12, 18)],
            [0, 35, 71, 107],
        )
        self.assertEqual(data.time_labels[35], "06:00:00")
        self.assertEqual(q234.interval_label(35), "6:00-6:10")

    def test_question_one_selected_times_and_blocks(self):
        selected = pd.read_csv(ROOT / "results" / "问题1_表1指定时段.csv")
        self.assertEqual(
            selected["时间段"].tolist(),
            ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"],
        )
        blocks = pd.read_csv(ROOT / "results" / "问题1_表2充放电汇总.csv")
        self.assertEqual(len(blocks), 6)
        self.assertAlmostEqual(float(blocks.loc[1, "放电量_kWh"]), 5947.419649152123, places=6)

    def test_result_workbooks_keep_official_interval_headers(self):
        for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
            workbook = load_workbook(ROOT / "results" / name, read_only=True, data_only=True)
            sheet = workbook["计划购电量"]
            if name == "result1.xlsx":
                self.assertEqual(sheet.cell(2, 1).value, "0:10-0:20")
                self.assertEqual(sheet.cell(145, 1).value, "0:00+1-0:10+1")
            else:
                self.assertEqual(sheet.cell(1, 2).value, "0:10-0:20")
                self.assertEqual(sheet.cell(1, 145).value, "0:00-0:10+1")

    def test_calendar_four_hour_summary_uses_previous_midnight_slot(self):
        frame = pd.read_csv(ROOT / "results" / "问题2至4_逐时段完整结果.csv", parse_dates=["日期"])
        frame = frame[frame["策略"] == "问题2_固定价"]
        previous_midnight = frame[
            (frame["日期"] == pd.Timestamp("2025-01-31")) & (frame["时间标签"] == "0:00+1")
        ]
        current = frame[frame["日期"] == pd.Timestamp("2025-02-01")].copy()
        current_minutes = current["时间标签"].map(q234._time_start_minutes)
        expected = pd.concat([previous_midnight, current[(current_minutes >= 10) & (current_minutes < 240)]])
        self.assertEqual(len(expected), 24)
        workbook = load_workbook(ROOT / "results" / "result2.xlsx", read_only=True, data_only=True)
        storage = workbook["充放电量"]
        self.assertAlmostEqual(storage.cell(2, 3).value, expected["充电量_kWh"].sum(), places=5)
        self.assertAlmostEqual(storage.cell(2, 4).value, expected["放电量_kWh"].sum(), places=5)


if __name__ == "__main__":
    unittest.main()
