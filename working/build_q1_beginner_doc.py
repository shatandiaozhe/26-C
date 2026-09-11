from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Sunuo\Desktop\26-C-master")
OUT = ROOT / "第一问零基础完整解释.docx"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def add_text(p, text, bold=False, color=None, size=None, font='宋体'):
    r=p.add_run(text); r.bold=bold
    if size: r.font.size=Pt(size)
    r.font.name=font; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
    if color: r.font.color.rgb=RGBColor(*color)
    return r

def equation(doc, text, note=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(5)
    if note:
        p.paragraph_format.keep_with_next = True
    r=add_text(p,text,size=11,font='Cambria Math'); r.italic=True
    if note:
        n=doc.add_paragraph(note); n.style='Caption'; n.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_table(doc, headers, rows, widths=None, numeric_cols=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False; table.style='Table Grid'
    hdr=table.rows[0]; set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; set_cell_shading(c,'234E70'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(p,str(h),bold=True,color=(255,255,255),size=9.5)
        set_cell_margins(c)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c)
            if ri%2: set_cell_shading(c,'F2F6FA')
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if numeric_cols and i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p,str(val),size=9.2)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); keep_with_next(p); return p

def body(doc,text,bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0.74); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.35
    if bold_lead and text.startswith(bold_lead):
        add_text(p,bold_lead,bold=True); add_text(p,text[len(bold_lead):])
    else: add_text(p,text)
    return p

def bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); add_text(p,text); return p

def numbered(doc,text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); add_text(p,text); return p

def add_figure(doc, rel, caption, width=15.3):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/rel),width=Cm(width))
    c=doc.add_paragraph(caption); c.style='Caption'; c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(8)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.35); sec.right_margin=Cm(2.35)
styles=doc.styles
normal=styles['Normal']; normal.font.name='宋体'; normal.font.size=Pt(10.5); normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
for name,size in [('Title',24),('Subtitle',12),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
    s=styles[name]; s.font.name='黑体' if name!='Subtitle' else '宋体'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0); s.font.bold=name!='Subtitle'; s._element.rPr.rFonts.set(qn('w:eastAsia'),s.font.name)
styles['Heading 1'].paragraph_format.space_before=Pt(14); styles['Heading 1'].paragraph_format.space_after=Pt(7)
styles['Heading 2'].paragraph_format.space_before=Pt(10); styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Caption'].font.name='宋体'; styles['Caption'].font.size=Pt(9); styles['Caption']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
# Word 的内置 Title 样式可能自带底部边框。解释文档只用留白区分标题。
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

# cover
p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
add_text(p,'第一问零基础完整解释',bold=True,size=24,font='黑体')
direct_border = p._p.get_or_add_pPr().find(qn('w:pBdr'))
if direct_border is not None:
    p._p.get_or_add_pPr().remove(direct_border)
p=doc.add_paragraph(); p.style='Subtitle'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
add_text(p,'微网购电与储能联合调度模型',size=14,font='宋体')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(30)
add_text(p,'从题意理解 变量定义 公式推导 到真实结果验证',size=12,font='宋体')
doc.add_paragraph('')
add_table(doc,['一眼看懂','本题中的含义'],[
    ['要决定什么','每 10 分钟买多少电 储能充多少 放多少 弃多少光'],
    ['必须保证什么','负荷始终得到满足 储能不越界 同时不能充放电 日末回到初始电量'],
    ['优化目标','全天向外部电网购电的费用最小'],
    ['最终结果','购电费 35126.95 元 比无储能少 12925.10 元 降幅 26.90%'],
], widths=[4.0,11.0])
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(55)
add_text(p,'适合第一次接触优化模型的读者',color=(70,70,70),size=10)
doc.add_page_break()

heading(doc,'阅读路线',1)
body(doc,'这份文档从生活化直觉开始，再把直觉翻译成变量和公式，最后回到程序输出和真实数据。建议第一次阅读按顺序看完。读到公式时，先看公式下方的中文解释，再回头辨认符号。')
add_table(doc,['部分','你会弄懂什么'],[
 ['一 题目到底在做什么','为什么这是一个按时间安排买电和用电池的问题'],
 ['二 原始数据怎样进入模型','10 分钟数据如何从千瓦换算为千瓦时'],
 ['三 参数和变量','哪些数是已知的 哪些数由求解器决定'],
 ['四 核心公式','供需平衡 储能递推 边界 互斥和目标函数'],
 ['五 求解器怎样找答案','为什么使用混合整数线性规划和两阶段优化'],
 ['六 用真实时段手算','低价充电 光伏富余充电 高价放电的具体数字'],
 ['七 全天结果和验证','最优成本 节省比例 灵敏度和残差意味着什么'],
 ['八 常见错误和复现','避免单位 效率 时段错位等问题'],
], widths=[5.0,10.0])
heading(doc,'先记住最终结论',1)
body(doc,'模型把储能看成一个受容量和效率限制的能量搬运工具。它在电价较低或光伏富余时充电，在电价较高且光伏不足时放电。由于日初和日末储电量都固定为 6000 kWh，节省不能靠把电池在一天结束时掏空，而只能来自合理的跨时段转移。')
add_table(doc,['指标','无储能','优化后','变化'],[
 ['全天购电量 kWh','61789.9354','59482.6990','减少 2307.2364'],
 ['全天购电费 元','48052.0466','35126.9486','节省 12925.0980'],
 ['成本降幅','0','26.8981%','约 26.90%'],
], widths=[4.1,3.4,3.4,4.1], numeric_cols={1,2,3})

heading(doc,'一 题目到底在做什么',1)
heading(doc,'1.1 把微网想成一间需要全天供电的房子',2)
body(doc,'小区负荷相当于每个时段必须完成的用电任务。光伏是当时能够免费使用的电，外部电网是随时可以买电的供应商，储能则像一个有容量上限的充电宝。不同之处在于，电价每 10 分钟可能不同，充电和放电还会损失一部分能量。')
body(doc,'因此模型每天都要回答同一组问题：当前时段的光伏先满足多少负荷，是否额外从电网买电给储能充电，储能是否应该放电减少高价购电，以及光伏太多且电池装不下时需要弃掉多少。')
heading(doc,'1.2 题目的硬条件',2)
for t in ['每个时段都必须满足负荷，不允许缺电。','储能电量必须一直处于 1200 到 10800 kWh。','充电和放电功率都不能超过 5000 kW。','同一时段不能既充电又放电。','不允许把电倒卖给外部电网，但允许弃光。','0:00 为 6000 kWh，24:00 也必须回到 6000 kWh。']:
    bullet(doc,t)
heading(doc,'1.3 为什么不能只看当前时段',2)
body(doc,'如果只看眼前，低价时段会只买刚好够负荷的电，但这样可能错过提前充电的机会。反过来，高价时段若随意放空电池，后面可能遇到更高电价。储能状态把 144 个时段连接起来，所以必须把整天一起优化。')

heading(doc,'二 原始数据怎样进入模型',1)
heading(doc,'2.1 一天被切成 144 个小格',2)
body(doc,'原数据每 10 分钟记录一次电价 负荷功率和光伏预测功率。一天有 24×6=144 个时段。设时段集合为 T={1,2,…,144}，单个时段长度为 10/60=1/6 小时。')
equation(doc,'Δt = 10 / 60 = 1 / 6 h','式 1  每个决策时段的长度')
heading(doc,'2.2 千瓦和千瓦时的区别',2)
body(doc,'千瓦 kW 表示此刻用电或发电的快慢，千瓦时 kWh 表示一段时间内的电量。储能容量和计划购电量都是 kWh，因此必须把附件中的功率乘以时段长度。')
equation(doc,'Lₜ = Pᴸₜ Δt = Pᴸₜ / 6     Rₜ = Pᴾⱽₜ Δt = Pᴾⱽₜ / 6','式 2  负荷功率和光伏功率转换为每时段电量')
body(doc,'例如第一条记录的负荷功率为 3439.8466 kW。这个功率持续 10 分钟，对应电量为 3439.8466÷6=573.3078 kWh。漏掉除以 6，会让购电量和储能变化全部放大 6 倍。')
heading(doc,'2.3 输入数据概况',2)
add_table(doc,['数据指标','实际数值','如何理解'],[
 ['时段数','144','覆盖完整一天'],['全天负荷电量','111024.8081 kWh','小区全天需要的总电量'],['全天光伏预测电量','55482.8357 kWh','当天可用的光伏总量'],['光伏富余时段','30 个','这些时段的光伏大于负荷'],['最低电价','0.3713 元/kWh','适合考虑充电'],['最高电价','1.3952 元/kWh','适合考虑放电'],['无储能购电费','48052.0466 元','用于比较优化是否真的省钱'],
], widths=[4.2,4.0,6.8], numeric_cols={1})
add_figure(doc,'figures/raw_q1_load_pv_price.png','图 1  原始负荷 光伏和电价的全天变化',15.5)

heading(doc,'三 已知参数和待求变量',1)
heading(doc,'3.1 已知参数',2)
body(doc,'参数是求解前已经知道的数，模型不能自行修改。带下标 t 的参数会随时段变化，不带下标的参数在全天保持不变。')
add_table(doc,['符号','含义','单位','数值或来源'],[
 ['t','时段序号','无','1 至 144'],['pₜ','第 t 时段电价','元/kWh','附件 1'],['Pᴸₜ','负荷功率','kW','附件 1'],['Pᴾⱽₜ','光伏预测功率','kW','附件 1'],['Lₜ','负荷电量','kWh','Pᴸₜ/6'],['Rₜ','可用光伏电量','kWh','Pᴾⱽₜ/6'],['Emin','储能安全下限','kWh','1200'],['Emax','储能安全上限','kWh','10800'],['E₀','初始储电量','kWh','6000'],['Pmax','最大充放电功率','kW','5000'],['qmax','每时段最大充放电量','kWh','5000/6=833.3333'],['ηch','充电效率','无','0.90'],['ηdis','放电效率','无','0.90'],
], widths=[2.0,5.3,2.7,5.0], numeric_cols={0,2,3})
heading(doc,'3.2 决策变量',2)
body(doc,'变量是求解器需要替我们决定的数。连续变量可以取小数，二元变量只能取 0 或 1。')
add_table(doc,['符号','中文含义','类型','单位','允许范围'],[
 ['gₜ','从外部电网购买的电量','连续','kWh','不小于 0'],['cₜ','交流母线送入储能的电量','连续','kWh','0 至 833.3333'],['dₜ','储能送回交流母线的电量','连续','kWh','0 至 833.3333'],['wₜ','没有被利用的光伏电量','连续','kWh','0 至 Rₜ'],['Eₜ','第 t 时段结束时的储电量','连续','kWh','1200 至 10800'],['uᶜʰₜ','该时段是否允许充电','二元','无','0 或 1'],['uᵈⁱˢₜ','该时段是否允许放电','二元','无','0 或 1'],
], widths=[2.1,5.5,2.0,2.0,3.4], numeric_cols={0,2,3,4})
body(doc,'特别注意 cₜ 和 dₜ 都按交流侧口径定义。cₜ=100 kWh 表示母线拿出 100 kWh 给电池充电，但效率为 0.9 时，电池只增加 90 kWh。dₜ=81 kWh 表示母线真正收到 81 kWh，此时电池内部要减少 81/0.9=90 kWh。')

heading(doc,'四 把题意翻译成公式',1)
heading(doc,'4.1 净负荷',2)
equation(doc,'Nₜ = Lₜ − Rₜ','式 3  不考虑储能时还差多少电')
body(doc,'Nₜ>0 说明光伏不够，需要电网或储能补足；Nₜ<0 说明光伏有富余，可以给储能充电或弃光。这个变量帮助理解情形，但不是必须由求解器决定的变量。')
heading(doc,'4.2 每个时段的供需平衡',2)
equation(doc,'gₜ + Rₜ + dₜ = Lₜ + cₜ + wₜ','式 4  进入母线的电量等于离开母线的电量')
body(doc,'左边是三个来源：电网购电 光伏发电和储能放电。右边是三个去向：小区负荷 储能充电和弃光。这个等式对 144 个时段逐一成立，因此不会凭空产生电，也不会让负荷缺电。')
heading(doc,'4.3 储能状态递推',2)
equation(doc,'Eₜ = Eₜ₋₁ + ηch cₜ − dₜ / ηdis','式 5  上一时段电量加有效充电量减内部放电消耗')
body(doc,'充电效率乘在 cₜ 前，是因为送进电池的 100 kWh 只能存下 90 kWh。放电量 dₜ 要除以放电效率，是因为要让母线得到 90 kWh，电池必须拿出 100 kWh。效率只在这条状态方程中出现，不能再在供需平衡式中重复计算。')
heading(doc,'4.4 容量和功率约束',2)
equation(doc,'1200 ≤ Eₜ ≤ 10800','式 6  储能始终在安全范围内')
equation(doc,'qmax = 5000 × 1/6 = 833.3333 kWh','式 7  10 分钟内最多充入或放出 833.3333 kWh')
body(doc,'储能总容量的安全可调区间为 10800−1200=9600 kWh。初始值 6000 kWh 恰好处于中点，因此向上和向下各有 4800 kWh 的空间。')
h=heading(doc,'4.5 充放电互斥',2); h.paragraph_format.page_break_before=True
equation(doc,'0 ≤ cₜ ≤ 833.3333 uᶜʰₜ     0 ≤ dₜ ≤ 833.3333 uᵈⁱˢₜ','式 8  状态开关控制充电和放电上限')
equation(doc,'uᶜʰₜ + uᵈⁱˢₜ ≤ 1     uᶜʰₜ,uᵈⁱˢₜ ∈ {0,1}','式 9  同一时段最多开启一种状态')
body(doc,'当充电开关为 0 时，第一条不等式强迫 cₜ=0；当放电开关为 0 时，dₜ=0。两个开关之和不超过 1，所以充电和放电不可能同时发生。二元变量也是模型名称中“整数”二字的来源。')
heading(doc,'4.6 禁止售电 允许弃光',2)
equation(doc,'gₜ ≥ 0     0 ≤ wₜ ≤ Rₜ','式 10  购电不能为负 弃光不能超过当时光伏')
body(doc,'如果允许 gₜ 为负，就等价于把多余电量卖给电网，但题目没有给出售电价格。加入弃光变量后，当光伏有富余且储能已满时，仍然能够保持能量平衡。')
heading(doc,'4.7 首末状态相等',2)
equation(doc,'E₀ = 6000     E₁₄₄ = 6000','式 11  日初和日末储电量相同')
body(doc,'这条约束保证比较公平。否则求解器会在最后几个高价时段把储能放到最低，把需要补回的电留给第二天，从而得到看似很低但不可持续的当天成本。')
heading(doc,'4.8 目标函数',2)
equation(doc,'Cgrid = Σₜ₌₁¹⁴⁴ pₜ gₜ     min Cgrid','式 12  把各时段购电量乘以电价后求和并取最小')
body(doc,'目标函数不直接要求购电量最少，而是要求费用最少。这允许模型在低价时多买一点，在高价时少买一点。只要低价充电的成本加上能量损失仍低于高价购电，就有套利价值。')
equation(doc,'pⱼ ηch ηdis > pᵢ','式 13  从低价时段 i 充电并在高价时段 j 放电的简化有利条件')
body(doc,'本数据最低价为 0.3713 元/kWh，最高价为 1.3952 元/kWh，往返效率为 0.9×0.9=0.81。计算 1.3952×0.81=1.1301，仍明显大于 0.3713，所以跨时段搬运电量在经济上可行。')

heading(doc,'五 为什么使用混合整数线性规划',1)
body(doc,'供需平衡 状态递推 成本和上下界都由变量的一次项组成，没有变量相乘或平方，因此属于线性关系。充放电开关只能取 0 或 1，因此整个问题是混合整数线性规划，简称 MILP。这里“混合”表示既有连续变量，也有整数变量。')
heading(doc,'5.1 求解器做了什么',2)
for t in ['先读取 144 个时段的电价 负荷和光伏。','为每个时段建立购电 充电 放电 弃光 SOC 和两个状态开关。','把全部公式写成矩阵形式交给 HiGHS 求解器。','求解器在满足所有约束的方案中寻找购电费最低的方案。','把答案回代到每条等式和不等式中做独立检查。']:
    numbered(doc,t)
heading(doc,'5.2 为什么还要两阶段优化',2)
body(doc,'不同调度方案可能拥有几乎相同的最低购电费。第一阶段只最小化购电费，得到 35126.948589 元。第二阶段把购电费限制在第一阶段最优值加 0.00001 元以内，再最小化全天充电量与放电量之和。')
equation(doc,'第一层  min Σpₜgₜ','式 14  先保证费用最低')
equation(doc,'第二层  min Σ(cₜ+dₜ)     且 Σpₜgₜ ≤ C*+10⁻⁵','式 15  在几乎不改变最低费用时减少无效循环')
add_table(doc,['阶段','购电费 元','总吞吐量 kWh','解释'],[
 ['第一阶段','35126.948589','37540.605698','只追求最低购电费'],['第二阶段','35126.948599','37540.605473','费用增加仅 0.00001 元 吞吐量更小'],
], widths=[3.2,3.3,3.5,5.0], numeric_cols={1,2})

heading(doc,'六 用真实时段手算',1)
heading(doc,'6.1 00:10 低价购电并满功率充电',2)
body(doc,'该时段电价为 0.4248 元/kWh，负荷电量 573.3078 kWh，光伏为 0。模型从电网购买 1406.6411 kWh，其中 573.3078 kWh 供负荷，833.3333 kWh 给储能充电。')
equation(doc,'1406.6411 + 0 + 0 = 573.3078 + 833.3333 + 0','式 16  供需平衡的数值回代')
equation(doc,'E₁ = 6000 + 0.9×833.3333 = 6750 kWh','式 17  充电损耗后电池实际增加 750 kWh')
heading(doc,'6.2 10:00 光伏富余直接充电',2)
body(doc,'该时段负荷为 983.7033 kWh，光伏为 1068.2657 kWh，富余 84.5624 kWh。模型不购电，把全部富余光伏送入储能，因此购电量和弃光量都为 0。')
equation(doc,'0 + 1068.2657 + 0 = 983.7033 + 84.5624 + 0','式 18  光伏先供负荷 多余部分充电')
heading(doc,'6.3 12:00 为什么光伏富余还要购电',2)
body(doc,'12:00 的光伏比负荷多 346.9304 kWh，但最大充电量为 833.3333 kWh。此时电价只有 0.4432 元/kWh，模型为了给后续高价时段准备能量，除用 346.9304 kWh 光伏富余充电外，还额外购电 486.4029 kWh，合计恰好满功率充电。')
equation(doc,'486.4029 + 1266.8406 = 919.9101 + 833.3333','式 19  低价购电与光伏富余共同给储能充电')
heading(doc,'6.4 18:10 高价时部分放电',2)
body(doc,'该时段净负荷为 641.1613 kWh，电价为 1.2561 元/kWh。模型让储能提供 109.2673 kWh，电网只需购买 531.8940 kWh。相比全部由电网供应，该时段少花约 137.25 元。')
equation(doc,'531.8940 + 45.9655 + 109.2673 = 687.1269','式 20  电网 光伏和储能共同满足负荷')
heading(doc,'6.5 20:40 最高价附近全由储能供电',2)
body(doc,'该时段电价达到全天最高值 1.3952 元/kWh，负荷电量为 715.6530 kWh，光伏为 0。模型不从电网买电，全部由储能放电满足负荷，时段结束时 SOC 正好到安全下限 1200 kWh。')
equation(doc,'0 + 0 + 715.6530 = 715.6530 + 0 + 0','式 21  高价时段的全部负荷由储能承担')
equation(doc,'E₁₂₄ = E₁₂₃ − 715.6530/0.9 = 1200 kWh','式 22  放电效率造成电池内部消耗更大')
heading(doc,'6.6 22:00 再充电为日末回归做准备',2)
body(doc,'22:00 电价回落到 0.4209 元/kWh。储能此时处于低位，模型购买 1393.7903 kWh，其中 560.4570 kWh 供负荷，833.3333 kWh 给储能充电。后续继续在低价时段补能，23:30 后 SOC 回到 6000 kWh。')

heading(doc,'七 全天最优结果怎样理解',1)
add_table(doc,['指标','结果','说明'],[
 ['求解状态','HiGHS Optimal','求解器找到满足精度要求的最优解'],['全天计划购电量','59482.6990 kWh','包含负荷用电和储能损耗所需购电'],['全天计划购电费','35126.9486 元','模型的首要目标值'],['节省金额','12925.0980 元','无储能成本减优化成本'],['成本降幅','26.8981%','节省金额除以无储能成本'],['全天充电量','20740.6660 kWh','交流侧送入储能的总电量'],['全天放电量','16799.9395 kWh','储能送回交流侧的总电量'],['弃光量','0 kWh','所有预测光伏均被负荷或储能利用'],['SOC 最低与最高','1200 至 10800 kWh','恰好触及安全边界但不越界'],['24:00 SOC','6000 kWh','满足日末回归'],
], widths=[4.2,4.2,6.6], numeric_cols={1})
add_figure(doc,'figures/process_q1_soc_actions.png','图 2  储能状态与充放电动作',15.5)
body(doc,'图中 SOC 在低价充电时上升，在高价放电时下降，并始终处于 1200 到 10800 kWh。充电总量 20740.6660 kWh 经过 0.81 的往返效率后可对应放电 16799.9395 kWh，数值上 20740.6660×0.81≈16799.9395。')
add_figure(doc,'figures/result_q1_cost_comparison.png','图 3  无储能方案与优化方案的全天购电费',13.2)

heading(doc,'八 怎样证明结果可信',1)
heading(doc,'8.1 逐时段回代',2)
body(doc,'把求解器输出重新代回供需平衡和储能递推公式。最大供需平衡残差为 2.274×10⁻¹³ kWh，最大 SOC 递推残差为 8.313×10⁻¹³ kWh，远小于检验容差 10⁻⁵ kWh。这些极小值来自计算机浮点数舍入，可视为数值上等于 0。')
add_table(doc,['检验项目','实测结果','判据','结论'],[
 ['供需平衡最大残差','2.274×10⁻¹³ kWh','≤10⁻⁵ kWh','通过'],['SOC 递推最大残差','8.313×10⁻¹³ kWh','≤10⁻⁵ kWh','通过'],['SOC 范围','1200 至 10800 kWh','在安全区间内','通过'],['日末偏差','0 kWh','≤10⁻⁵ kWh','通过'],['同时充放电时段','0 个','必须为 0','通过'],['全天能量恒等式残差','5.457×10⁻¹² kWh','≤10⁻⁵ kWh','通过'],
], widths=[4.4,4.1,3.5,3.0], numeric_cols={1,2,3})
heading(doc,'8.2 和无储能基准比较',2)
body(doc,'若完全不用储能，每个时段只在光伏不足时购电，费用为 48052.0466 元。优化结果为 35126.9486 元，确实更低。这个比较证明储能策略有经济收益，但仅凭它还不能证明全局最优。')
heading(doc,'8.3 和连续松弛下界比较',2)
body(doc,'把充放电互斥的二元限制暂时放宽，会得到一个更容易但可行域更大的线性问题，其最低费用是原问题不可能突破的下界 35126.948589 元。最终结果只比下界高 0.00001 元，恰好对应第二阶段允许的成本容差，因此可认为在给定数值精度内达到全局最优。')

heading(doc,'九 效率变化会怎样',1)
body(doc,'主方案把充电效率和放电效率都取 0.90，往返效率为 0.81。效率越低，搬运同样多的电就需要额外购买更多电量，因此成本会上升。若题目中的 90% 被解释为总往返效率，应令两个单程效率都等于 √0.9≈0.948683。')
add_table(doc,['单程效率','往返效率','最优购电费 元','全天购电量 kWh'],[
 ['0.80','0.64','38223.6516','64053.8678'],['0.85','0.7225','36603.4065','61728.9749'],['0.90','0.81','35126.9486','59482.6990'],['0.948683','0.90','33801.4956','57526.2435'],['0.95','0.9025','33767.0326','57472.6431'],['1.00','1.00','32377.0883','55541.9724'],
], widths=[3.3,3.3,4.2,4.2], numeric_cols={0,1,2,3})
add_figure(doc,'figures/result_q1_efficiency_sensitivity.png','图 4  单程效率变化对最优购电费的影响',14.5)
body(doc,'这项分析说明模型结论的方向稳定：无论采用哪种合理效率，储能都能利用电价差降低购电费；但具体费用会随效率口径变化，因此正式答题时应明确写出效率解释。')

heading(doc,'十 最容易犯的错误',1)
add_table(doc,['错误','会造成什么问题','正确做法'],[
 ['把 kW 直接当 kWh','所有电量与 SOC 变化放大 6 倍','每 10 分钟功率都乘 1/6 小时'],['在两条方程都乘效率','效率被重复计算 能量凭空减少','效率只进入 SOC 递推式'],['允许购电量为负','模型偷偷把电卖给电网','设置 gₜ≥0 并单独设置弃光变量'],['没有充放电互斥','可能出现同一时段同时充放电','使用两个二元开关及和不超过 1'],['删除日末 SOC 约束','模型在最后透支电池 费用虚低','强制 E₁₄₄=E₀=6000'],['按时间文本直接连接','00:10 与跨日标签可能整体错位','内部统一用 1 至 144 的时段序号'],['只看成本不验约束','可能接受物理上不可执行的结果','回代平衡 SOC 边界 互斥和终端条件'],
], widths=[3.8,5.4,5.8])

heading(doc,'十一 从模型到结果文件',1)
add_table(doc,['结果文件中的内容','对应模型量'],[
 ['每 10 分钟计划购电量','直接输出 gₜ'],['题目指定时段购电量','按时段序号读取对应 gₜ'],['全天购电量','Σgₜ'],['全天购电费','Σpₜgₜ'],['每 4 小时充电量','时间块内 Σcₜ'],['每 4 小时放电量','时间块内 Σdₜ'],['0:00 储电量','E₀=6000'],['24:00 储电量','E₁₄₄=6000'],
], widths=[7.2,7.8])
heading(doc,'11.1 题目指定时段购电量',2)
add_table(doc,['时间段','购电量 kWh'],[
 ['10:00 至 10:10','0.0000'],['12:00 至 12:10','486.4029'],['14:00 至 14:10','0.0000'],['16:00 至 16:10','394.9315'],['18:00 至 18:10','636.9826'],['20:00 至 20:10','0.0000'],
], widths=[8.0,7.0], numeric_cols={1})
heading(doc,'11.2 每四小时充放电汇总',2)
add_table(doc,['时间段','充电量 kWh','放电量 kWh'],[
 ['0:00 至 4:00','4500.0000','0.0000'],['4:00 至 8:00','833.3333','5947.4196'],['8:00 至 12:00','3954.6308','2121.4184'],['12:00 至 16:00','6119.3685','91.1014'],['16:00 至 20:00','0.0000','5068.6869'],['20:00 至 24:00','5333.3333','3571.3132'],
], widths=[6.0,4.5,4.5], numeric_cols={1,2})

heading(doc,'十二 一页复习',1)
body(doc,'第一问的本质是：在 144 个相互连接的时段中，安排购电和储能动作，使负荷全部得到满足，同时让全天购电费最低。')
for t in ['先把负荷和光伏从 kW 转为 kWh。','用供需平衡保证每个时段的电有来源也有去向。','用 SOC 递推把前后时段连接，并正确处理 0.9 的充放电效率。','用容量 功率 互斥 禁止售电和首末相等约束保证方案可执行。','用 Σpₜgₜ 作为首要目标，再用第二阶段减少无效充放电。','最终购电费 35126.95 元，比无储能节省 12925.10 元，降幅 26.90%。','所有关键残差低于 10⁻¹² kWh 量级，同时充放电时段为 0，日末 SOC 回到 6000 kWh。']:
    bullet(doc,t)
heading(doc,'结语',1)
body(doc,'如果只记一个画面，可以把储能理解为在时间轴上搬运电量的仓库：低价或光伏富余时入库，高价且光伏不足时出库。公式的作用，是给这个仓库加上真实的容量 效率 功率和操作规则；求解器的作用，是在所有合法搬运方案中找到总费用最低的那一个。')

# footer page number field
for section in doc.sections:
    footer=section.footer; p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_text(p,'第 ',size=9,color=(90,90,90))
    run=p.add_run(); fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='PAGE'; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.extend([fldChar1,instr,fldChar2])
    add_text(p,' 页',size=9,color=(90,90,90))

doc.core_properties.title='第一问零基础完整解释'
doc.core_properties.subject='微网购电与储能联合调度模型'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)
